#!/usr/bin/env python3
"""Laeb RHR-ist hanke kehtivad alusdokumendid (zip) ja teisendab tekstiks.
Kasutus: python3 rhr_docs.py <RHR procurement id, nt 10682825> [väljundkaust]
Avalik API, sisselogimist ei vaja. PDF->txt vajab `pdftotext` (poppler), DOCX->txt `python-docx`."""
import sys, os, json, zipfile, io, subprocess, urllib.request
B='https://riigihanked.riik.ee'
pid=sys.argv[1]; out=sys.argv[2] if len(sys.argv)>2 else f'rhr_{pid}'
os.makedirs(out+'/docs',exist_ok=True); os.makedirs(out+'/txt',exist_ok=True)
rel=json.load(urllib.request.urlopen(f'{B}/rhr/api/public/v1/procurement/{pid}/documents-temp-url',timeout=60))['value']
z=zipfile.ZipFile(io.BytesIO(urllib.request.urlopen(B+rel,timeout=300).read())); z.extractall(out+'/docs')
for f in sorted(os.listdir(out+'/docs')):
    src=f'{out}/docs/{f}'; base,ext=os.path.splitext(f); dst=f'{out}/txt/{base}.txt'
    try:
        if ext.lower()=='.pdf': subprocess.run(['pdftotext','-layout',src,dst],check=True)
        elif ext.lower()=='.docx':
            import docx; d=docx.Document(src)
            open(dst,'w').write('\n'.join(p.text for p in d.paragraphs)+'\n'+'\n'.join(' | '.join(c.text for c in r.cells) for t in d.tables for r in t.rows))
    except Exception as e: print('!',f,e)
    print(f)
print('->',out)
